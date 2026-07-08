#coding = utf-8
"""将表格数据集导出为 Excel HttpResponse。"""

from __future__ import annotations

import io
import json
import re
from typing import Any, Dict, List, Optional

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from assistant.export_fields import build_field_glossary_for_datasets, get_field_meta


def friendly_sheet_title(name: str, idx: int, part: int = 1) -> str:
    base = (name or f"表{idx}").strip()
    replace_map = {
        "cashier_trade_summary": "交易汇总",
        "cashier_paymode_summary": "付款方式",
        "cashier_item_summary": "项目明细",
        "store_dimension_sales_summary": "管理维度营业",
        "cashier_archievement_summary": "员工业绩",
        "vip_maintenance_summary": "会员维护",
        "vip_top_cash_consumption": "会员现金消费TOP",
        "vip_top_card_consumption": "会员卡付消费TOP",
        "vip_top_send_consumption": "会员赠送消费TOP",
        "vip_top_service_consumption": "会员服务消费TOP",
        "vip_top_goods_consumption": "会员商品消费TOP",
        "vip_sleeping_alert": "沉睡会员预警",
        "vip_sleeping_alert_sleeping_vips": "沉睡会员明细",
        "panel_store_dimension_sales": "管理维度营业",
        "panel_vip_sleeping_alert": "沉睡会员预警",
        "batch_store_vips": "门店会员列表",
        "time_dimension_summary": "时间趋势",
        "list_recent_expvstoll": "交易流水",
        "readonly_sql": "SQL查询结果",
        "get_vip_detail": "会员档案",
        "vip_profile": "客户画像",
        "vip_churn_risk": "流失风险",
        "search_vips": "会员搜索",
    }
    base = replace_map.get(base, base)
    base = re.sub(r"[^0-9A-Za-z一-鿿_\-]", "_", base)
    if not base:
        base = f"表{idx}"
    if part > 1:
        base = f"{base}_{part}"
    return base[:31]


def detect_column_type(col: str) -> str:
    c = (col or "").lower()
    if c in {"vsdate", "cdate", "date", "indate", "opendate", "issuedate"} or c.endswith("date"):
        return "date"
    if c in {"create_time", "updated_at", "last_modified", "time"} or c.endswith("_time"):
        return "datetime"
    if any(k in c for k in ["amount", "mount", "money", "price", "cost", "total", "avg", "arch"]):
        return "money"
    if any(k in c for k in ["qty", "count", "num", "times", "days"]):
        return "number"
    return "text"


def _try_parse_yyyymmdd(value: Any):
    if value is None:
        return None
    s = str(value).strip().replace("-", "")
    if len(s) == 8 and s.isdigit():
        try:
            from datetime import datetime as _dt
            return _dt.strptime(s, "%Y%m%d").date()
        except ValueError:
            return None
    return None


def _try_parse_datetime(value: Any):
    if value is None:
        return None
    s = str(value).strip()
    from datetime import datetime as _dt
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M"):
        try:
            return _dt.strptime(s, fmt)
        except ValueError:
            continue
    return None


def coerce_cell_value(v: Any, col_type: str):
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    if col_type == "date":
        d = _try_parse_yyyymmdd(v)
        if d is not None:
            return d
    if col_type == "datetime":
        dt = _try_parse_datetime(v)
        if dt is not None:
            return dt
    if col_type in {"money", "number"}:
        if v is None or v == "":
            return None
        try:
            return float(v)
        except (TypeError, ValueError):
            return v
    return v


def _append_glossary_sheet(wb: Workbook, glossary: List[Dict[str, str]]) -> None:
    if not glossary:
        return
    ws = wb.create_sheet(title="字段说明"[:31])
    header = ["数据集", "字段名", "中文名称", "字段说明"]
    ws.append(header)
    for row in glossary:
        ws.append(
            [
                row.get("dataset") or "",
                row.get("field") or "",
                row.get("label") or "",
                row.get("description") or "",
            ]
        )


def build_xlsx_http_response(
    datasets: List[Dict[str, Any]],
    *,
    filename: str = "assistant-export.xlsx",
    max_rows_per_sheet: int = 50000,
    field_glossary: Optional[List[Dict[str, str]]] = None,
) -> HttpResponse:
    wb = Workbook(write_only=True)
    glossary = field_glossary if field_glossary is not None else build_field_glossary_for_datasets(datasets)
    _append_glossary_sheet(wb, glossary)

    ds_idx = 0
    for ds in datasets:
        rows: List[Dict[str, Any]] = ds.get("rows") or []
        if not rows:
            continue
        ds_idx += 1
        ds_name = str(ds.get("name") or "sheet")
        cols: List[str] = []
        seen = set()
        for r in rows:
            for k in r.keys():
                k2 = str(k)
                if k2 not in seen:
                    seen.add(k2)
                    cols.append(k2)
        if not cols:
            continue
        col_types = {c: detect_column_type(c) for c in cols}
        col_labels = {c: get_field_meta(c, ds_name).get("label") or c for c in cols}
        total = len(rows)
        part = 1
        start_pos = 0
        while start_pos < total:
            end_pos = min(start_pos + max_rows_per_sheet, total)
            part_rows = rows[start_pos:end_pos]
            title = friendly_sheet_title(ds_name, ds_idx, part=part)
            ws = wb.create_sheet(title=title)
            label_header = [WriteOnlyCell(ws, value=col_labels.get(c, c)) for c in cols]
            ws.append(label_header)
            key_header = [WriteOnlyCell(ws, value=c) for c in cols]
            ws.append(key_header)
            for r in part_rows:
                line = []
                for c in cols:
                    v = coerce_cell_value(r.get(c), col_types.get(c, "text"))
                    cell = WriteOnlyCell(ws, value=v)
                    typ = col_types.get(c, "text")
                    if typ == "money":
                        cell.number_format = "#,##0.00"
                    elif typ == "number":
                        cell.number_format = "0.00"
                    elif typ == "date":
                        cell.number_format = "yyyy-mm-dd"
                    elif typ == "datetime":
                        cell.number_format = "yyyy-mm-dd hh:mm:ss"
                    line.append(cell)
                ws.append(line)
            start_pos = end_pos
            part += 1
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(
        buf.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp
