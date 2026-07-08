# -*- coding: utf-8 -*-
"""Django admin: export selected rows to .xlsx using list_display (openpyxl)."""
import re
from io import BytesIO
from datetime import date, datetime
from decimal import Decimal
from urllib.parse import quote

from django.http import HttpResponse
from django.utils import timezone
from django.utils.encoding import force_str


def _cell_value(modeladmin, request, obj, field_name):
    """Match changelist: ModelAdmin list_display (methods + model fields)."""
    from django.db.models import FileField, ImageField

    if field_name in ("action_checkbox", "pk"):
        return ""

    if hasattr(modeladmin, field_name):
        fn = getattr(modeladmin, field_name, None)
        if callable(fn):
            try:
                return fn(obj)
            except TypeError:
                return ""

    try:
        f = modeladmin.model._meta.get_field(field_name)
        if f.many_to_many:
            return ", ".join(
                force_str(x) for x in getattr(obj, f.name).all()[:200]
            )
        v = f.value_from_object(obj)
        if f.is_relation:
            if v is None:
                return ""
            return force_str(v)
        if v is None:
            return ""
        if isinstance(f, (FileField, ImageField)) and v:
            return v.name
        if isinstance(v, (datetime, date)):
            return v.isoformat()
        if isinstance(v, Decimal):
            return float(v)
        if isinstance(v, (bytes, memoryview)):
            return "<binary>"
        return v
    except Exception:
        pass

    try:
        v = getattr(obj, field_name)
        if callable(v):
            v = v()
        if v is None:
            return ""
        if isinstance(v, (datetime, date)):
            return v.isoformat()
        return v
    except Exception:
        return ""


def _stringify(val):
    if val is None:
        return ""
    if isinstance(val, (datetime, date)):
        return val.isoformat()
    if isinstance(val, Decimal):
        return str(val)
    return force_str(val)


def modeladmin_export_xlsx(modeladmin, request, queryset):
    """
    Build .xlsx for queryset using changelist's list_display columns.
    Returns HttpResponse, or None if queryset empty.
    """
    from openpyxl import Workbook

    if not queryset.exists():
        return None

    list_display = [
        n
        for n in modeladmin.get_list_display(request)
        if n and n not in ("action_checkbox",)
    ]
    if not list_display:
        list_display = [f.name for f in modeladmin.model._meta.concrete_fields]

    wb = Workbook()
    ws = wb.active
    ws.title = "export"
    model = modeladmin.model

    for col, name in enumerate(list_display, 1):
        header = name
        try:
            f = model._meta.get_field(name)
            header = str(f.verbose_name)
        except Exception:
            pass
        ws.cell(row=1, column=col, value=_stringify(header))

    for row, obj in enumerate(queryset, 2):
        for col, name in enumerate(list_display, 1):
            raw = _cell_value(modeladmin, request, obj, name)
            ws.cell(row=row, column=col, value=_stringify(raw))

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_name = re.sub(r"[^\w\-.]+", "_", model._meta.model_name)[:50]
    filename = "%s_%s.xlsx" % (safe_name, timezone.now().strftime("%Y%m%d_%H%M%S"))
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = "attachment; filename=%s" % quote(
        filename
    )  # filename as ASCII/quoted; weixin browsers OK

    return response
